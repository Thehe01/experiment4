"""Synchronize Chapter 3 prose with the executable annotation boundary."""

from __future__ import annotations

import os
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[3]
TARGET = ROOT / "融合攻击模型的安全漏洞知识表示方法_修改稿.docx"
TEMP = ROOT / ".tmp" / "chapter3_boundary_sync.docx"


APPENDS = {
    "据此，完整本体在实例化过程中分为文本抽取层和映射补全层。": (
        "在文本提及层，Configuration 的表面跨度采用最小但可独立识别的产品或组件名称，"
        "普通版本、补丁和构建号不进入跨度；完整 CPE URI 以及唯一规范化所需的约定产品"
        "代际标识除外。Weakness 的描述性短语只有在同一事实块明示唯一对应的 CWE 时"
        "方可进入正式结果，否则只保留为复核候选。"
    ),
    "上述分层同时构成第4章实验的任务边界。": (
        "后续抽取方法只能在该边界内改进实体定位和关系判定，不得通过提示优化改变"
        "实体定义、关系方向、表面跨度政策或直接证据条件。"
    ),
    "资产层包含 Configuration 类": (
        "本体中的 Configuration 可以通过规范化标识和属性记录版本范围；文本中的"
        "Configuration 表面跨度与该规范化配置分开保存，普通版本号、补丁号和构建号"
        "不并入产品提及。"
    ),
    "exploited_by 关系连接具体漏洞与 ATT&CK 攻击技术": (
        "例如，文本表述攻击者利用某 CVE 后运行远程访问木马并建立命令与控制通道时，"
        "该证据只能支持后续攻击活动，不能据此把命令与控制技术作为该 CVE 的"
        "exploited_by 尾实体。类似地，ATT&CK 技术标签与 CVE 位于同一表格行，也仍需"
        "核对行为描述是否表明该技术直接承担漏洞利用，而不能只凭结构邻接建立关系。"
    ),
}

OLD_INSTANTIATES = (
    "在文本未给出 CWE 编号但明确描述漏洞成因时，第4章可以将相应描述归一化为候选 "
    "Weakness，并保留文本依据。"
)
NEW_INSTANTIATES = (
    "在文本未给出 CWE 编号但明确描述漏洞成因时，第4章只将相应描述保留为候选 "
    "Weakness；只有同一事实块出现唯一对应的明示 CWE，或后续另行冻结并版本化受控"
    "词表后，才可晋级正式实体及 instantiates 关系。"
)

OLD_LOG4J_CONFIGURATION = (
    "在本体中，受影响的 Log4j2 产品及版本表示为 Configuration，CVE-2021-44228 "
    "表示为 Vulnerability。"
)
NEW_LOG4J_CONFIGURATION = (
    "在本体中，受影响的 Log4j2 产品表示为 Configuration，2.14.1 等版本信息作为 "
    "version 属性与实体分开保存，CVE-2021-44228 表示为 Vulnerability。"
)

FIGURE_BOUNDARY_NOTE = (
    "图中 Configuration 框内的“Log4j2”是实体名称，“2.14.1”是独立的 version "
    "属性展示，并不表示第4章应将二者合并为文本提及跨度。"
)


def main() -> None:
    if not TARGET.is_file():
        raise FileNotFoundError(TARGET)
    document = Document(TARGET)
    append_hits = {prefix: 0 for prefix in APPENDS}
    replacement_hits = 0
    log4j_configuration_hits = 0
    figure_note_hits = 0

    for paragraph in document.paragraphs:
        text = paragraph.text
        if OLD_INSTANTIATES in text:
            paragraph.text = text.replace(OLD_INSTANTIATES, NEW_INSTANTIATES)
            replacement_hits += 1
            text = paragraph.text
        elif NEW_INSTANTIATES in text:
            replacement_hits += 1
        if OLD_LOG4J_CONFIGURATION in paragraph.text:
            paragraph.text = paragraph.text.replace(
                OLD_LOG4J_CONFIGURATION, NEW_LOG4J_CONFIGURATION
            )
            log4j_configuration_hits += 1
            text = paragraph.text
        elif NEW_LOG4J_CONFIGURATION in paragraph.text:
            log4j_configuration_hits += 1
        for prefix, addition in APPENDS.items():
            if text.startswith(prefix):
                append_hits[prefix] += 1
                if addition not in paragraph.text:
                    paragraph.add_run(addition)
        if "文本中的Configuration" in paragraph.text or "CVE 的exploited_by" in paragraph.text:
            paragraph.text = (
                paragraph.text
                .replace("文本中的Configuration", "文本中的 Configuration")
                .replace("CVE 的exploited_by", "CVE 的 exploited_by")
            )
        if paragraph.text.startswith("图3-3从上至下区分攻击过程事实与本体实例关系"):
            figure_note_hits += 1
            if FIGURE_BOUNDARY_NOTE not in paragraph.text:
                paragraph.add_run(FIGURE_BOUNDARY_NOTE)

    if replacement_hits != 1:
        raise RuntimeError(
            f"expected one instantiates paragraph replacement, got {replacement_hits}"
        )
    if log4j_configuration_hits != 1:
        raise RuntimeError(
            "expected one Log4j2 configuration paragraph replacement, "
            f"got {log4j_configuration_hits}"
        )
    if figure_note_hits != 1:
        raise RuntimeError(
            f"expected one figure 3-3 boundary note anchor, got {figure_note_hits}"
        )
    missing = [prefix for prefix, count in append_hits.items() if count != 1]
    if missing:
        raise RuntimeError(f"chapter paragraph anchors changed: {missing}")

    relation_table = document.tables[1]
    instantiates_basis = (
        "CVE、NVD 等条目中的 CWE 分类字段，或同一事实块中与漏洞成因描述唯一对应的"
        "明示 CWE"
    )
    if relation_table.cell(2, 0).text.strip() != "instantiates":
        raise RuntimeError("table 3-2 row layout changed")
    relation_table.cell(2, 4).text = instantiates_basis

    TEMP.parent.mkdir(parents=True, exist_ok=True)
    document.save(TEMP)
    os.replace(TEMP, TARGET)
    print(TARGET)


if __name__ == "__main__":
    main()
